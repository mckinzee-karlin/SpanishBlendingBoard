import json
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
d=json.load(open('data/data.json'));L=json.load(open('data/lists.json'));d['SENT']={k:[v] for k,v in L['sent'].items()}
doc=Document()
st=doc.styles['Normal'];st.font.name='Verdana';st.font.size=Pt(12)
st.element.rPr.rFonts.set(qn('w:eastAsia'),'Verdana')
for s in ['Heading 1','Heading 2','Title']:
    doc.styles[s].font.name='Verdana'
# language tag
rpr=st.element.get_or_add_rPr();l=OxmlElement('w:lang');l.set(qn('w:val'),'es-MX');rpr.append(l)
doc.core_properties.title='Letra del día: guion de la rutina'
doc.core_properties.language='es-MX'
doc.add_heading('Letra del día: guion de la rutina',0)
doc.add_paragraph('Rutina explícita de conocimiento del alfabeto para estudiantes que están aprendiendo las letras. Cada paso tiene tres rondas. En cada ronda, cambie lo que señala: la mayúscula, la minúscula o las dos. Sustituya [letra], [nombre], [sonido] y [palabra] por los datos de la tabla de la última sección.')
def sec(h,items):
    doc.add_heading(h,1)
    for kind,t in items:
        p=doc.add_paragraph(t,style='List Bullet' if kind=='b' else None)
        if kind=='s':
            p.paragraph_format.left_indent=Pt(18)
            for r in p.runs:r.italic=True
sec('1. Nombrarla',[('s','“Esta es la letra [nombre]. Esta es la [nombre] mayúscula y esta es la [nombre] minúscula.”'),('p','Señale cada letra. Pregunte: “¿Qué letra es esta?” Si el estudiante no responde, diga el nombre y pida que lo repita o lo señale con su método de respuesta.')])
sec('2. Su sonido',[('s','“Esta letra representa este sonido: /[sonido]/. Escuchen: /[sonido]/ como en [palabra].”'),('p','Muestre la imagen de la palabra ancla. Estire el primer sonido de la palabra. Pregunte: “¿Qué sonido representa esta letra?” La h no tiene sonido. Diga: “La h se escribe, pero no suena.”')])
sec('3. Encontrarla',[('s','“Ahora busquemos la letra [nombre]. Cuando la encuentren, digan el nombre y el sonido.”'),('p','La aplicación lee la oración y resalta cada palabra. Después el estudiante busca la letra en la oración y la toca. Cuando la encuentra, diga el nombre y el sonido. La marca grande aparece sola cuando se encuentran todas.')])
sec('Sílabas (consonantes)',[('p','Toque una sílaba de la fila inferior. Se abre una ventana con dos palabras e imágenes. Diga la palabra y señale la sílaba resaltada. Los estudiantes escuchan y miran la sílaba. No leen las palabras.')])
sec('4. Escribirla',[('s','“Voy a mostrarles cómo se escribe la letra [nombre]. Primero la mayúscula y después la minúscula. Mientras escribo, digo cada paso.”'),('p','Lea los pasos de la guía de trazo mientras traza. El estudiante traza en el aire, en papel, en una superficie o con un lápiz alternativo. Muestre cómo usar ese lápiz.')])
sec('Cómo responde el estudiante sin hablar',[('b','Señala la letra, la imagen o la opción con el dedo, la mano, un puntero o la cabeza.'),('b','Mira la opción cuando las tarjetas se presentan separadas y quietas.'),('b','Usa un dispositivo de comunicación (AAC), un teclado o un tablero de letras.'),('b','Si no puede señalar, use el barrido asistido por un compañero: señale o lea cada opción y el estudiante responde sí o no con su propio método.')])
doc.add_heading('Antes de usar',1)
for t in ['Consulte con el logopeda (SLP) o el especialista en AAC antes de pedir sonidos aislados a estudiantes con apraxia del habla, disartria o sin habla funcional.','Consulte con el especialista en discapacidad visual (TVI) para el tamaño, el contraste y el formato táctil de las tarjetas.','Consulte con el maestro de sordos o con audiología antes de usar los sonidos grabados con estudiantes con pérdida auditiva.','Las palabras ancla marcadas con * se tomaron de la lección de Bluebonnet que enseña la letra. Las demás son propuestas del adaptador y deben verificarse.']:
    doc.add_paragraph(t,style='List Bullet')
doc.add_heading('Tabla por letra',1)
tbl=doc.add_table(rows=1,cols=5);tbl.style='Table Grid'
for c,h in zip(tbl.rows[0].cells,['Letra','Nombre','Sonido','Palabra ancla','Oración para encontrarla']):
    c.text=h
    for r in c.paragraphs[0].runs:r.bold=True
trPr=tbl.rows[0]._tr.get_or_add_trPr();th=OxmlElement('w:tblHeader');th.set(qn('w:val'),'true');trPr.append(th)
VERIF={'m','j','h','k','x','ch'}
for t in d['LET']:
    snd='(ninguno)' if t=='h' else ('/s/ al inicio, /ks/ en medio' if t=='x' else '/'+t+'/')
    row=tbl.add_row().cells
    vals=[(t if t=='rr' else t[0].upper()+t[1:]+' '+t),d['NAMES'][t],snd,d['ANCH'][t]+(' *' if t in VERIF else ''),d['SENT'][t][0]]
    for c,v in zip(row,vals):c.text=v
doc.add_paragraph()
doc.add_paragraph('Fuentes: Explicit Alphabet Knowledge Instructional Routine (Jones et al., 2013), adaptada al español. Palabras ancla y guía de trazo: Bluebonnet Learning, Texas Education Agency, CC BY-NC-SA 4.0. Las oraciones y las palabras de las sílabas las escribió el adaptador y deben revisarse.')
doc.save('../docs/Letra_del_dia_guion.docx')

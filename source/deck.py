import json
from pptx import Presentation
from pptx.util import Emu,Pt
from pptx.dml.color import RGBColor
L=json.load(open('data/lists.json'))['sent']
EN={'a':'The spider walks on the branch.','b':'Beto kisses his baby.','c':'Cuca eats coconut.','ch':'Chucho drinks milk.','d':'Dina gives a candy.','e':'The elephant is enormous.','f':'Fifi the seal uses a photo.','g':'Gabo the cat keeps the eraser.','h':'Hilda eats ice cream.','i':'The iguana is on the island.','j':'Jacinta the giraffe jumps.','k':'Karen does karate.','l':'Lalo the parrot sees the moon.','ll':'The rain falls on the street.','m':'The monkey looks at the table.','n':'Nico the little boy sees a cloud.','ñ':'The girl Toña sees a piñata.','o':'Olga sees a bear.','p':'Pepe paints a duck.','qu':'Quique wants cheese.','r':'Rita sees a frog.','rr':'The dog runs to the car.','s':'Sofía takes her chair out to the sun.','t':'Teo has a taco.','u':'Úrsula eats grapes.','v':'Vera sees a cow.','w':'Wanda eats a waffle.','x':'The taxi carries a xylophone.','y':'Yola uses a yo-yo.','z':'The fox uses a shoe.'}
prs=Presentation();prs.slide_width=Emu(12192000);prs.slide_height=Emu(6858000)
blank=prs.slide_layouts[6]
for k,es in L.items():
    s=prs.slides.add_slide(blank)
    tb=s.shapes.add_textbox(Emu(300000),Emu(6250000),Emu(11600000),Emu(500000));tb.name='Sentence and translation'
    tf=tb.text_frame;tf.word_wrap=True
    for i,(txt,lang) in enumerate([(es,'es-MX'),(EN[k],'en-US')]):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        r=p.add_run();r.text=txt;r.font.size=Pt(11);r.font.name='Verdana';r.font.color.rgb=RGBColor(0x59,0x59,0x59)
        r.font._element.set('lang',lang)
prs.core_properties.title='Letra del día: sentence card backgrounds'
prs.save('../docs/Letra_del_dia_oraciones_en_blanco.pptx')
print(len(prs.slides))

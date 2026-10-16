import json
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
d=json.load(open('data/data.json'));L=json.load(open('data/lists.json'))['sent']
EN={'a':'The spider walks on the branch.','b':'Beto kisses his baby.','c':'Cuca eats coconut.','ch':'Chucho drinks milk.','d':'Dina gives a candy.','e':'The elephant is enormous.','f':'Fifi the seal uses a photo.','g':'Gabo the cat keeps the eraser.','h':'Hilda eats ice cream.','i':'The iguana is on the island.','j':'Jacinta the giraffe jumps.','k':'Karen does karate.','l':'Lalo the parrot sees the moon.','ll':'The rain falls on the street.','m':'The monkey looks at the table.','n':'Nico the little boy sees a cloud.','ñ':'The girl Toña sees a piñata.','o':'Olga sees a bear.','p':'Pepe paints a duck.','qu':'Quique wants cheese.','r':'Rita sees a frog.','rr':'The dog runs to the car.','s':'Sofía takes her chair out to the sun.','t':'Teo has a taco.','u':'Úrsula eats grapes.','v':'Vera sees a cow.','w':'Wanda eats a waffle.','x':'The taxi carries a xylophone.','y':'Yola uses a yo-yo.','z':'The fox uses a shoe.'}
AEN={'a':'spider','b':'robe','c':'house','ch':'chocolate','d':'finger','e':'elephant','f':'seal','g':'cat','h':'ice cream','i':'iguana','j':'giraffe','k':'kiwi','l':'parrot','ll':'key','m':'monkey','n':'little boy','ñ':'pineapple','o':'bear','p':'duck','qu':'cheese','r':'frog','rr':'dog','s':'toad','t':'shark','u':'fingernail','v':'cow','w':'wifi','x':'xylophone','y':'yo-yo','z':'fox'}
SND={'h':'none (silent letter)','x':'/s/ at the start of a word; /ks/ in the middle','c':'/k/ (before a, o, u); /s/ (before e, i)','g':'/g/ (before a, o, u); /j/ (before e, i)','rr':'trilled /rr/ (between vowels only)','r':'/r/ (tap in the middle of a word; trill at the start)','qu':'/k/ (qu is followed by e or i)','ch':'/ch/','ll':'/ll/ (same sound as y in many dialects)','ñ':'/ñ/','z':'/s/ (Latin American Spanish)','v':'/b/ (Mexican Spanish)','y':'/y/','w':'/w/'}
doc=Document()
st=doc.styles['Normal'];st.font.name='Verdana';st.font.size=Pt(11)
st.element.rPr.rFonts.set(qn('w:eastAsia'),'Verdana')
for s in ['Heading 1','Heading 2','Title']:doc.styles[s].font.name='Verdana'
rpr=st.element.get_or_add_rPr();l=OxmlElement('w:lang');l.set(qn('w:val'),'en-US');rpr.append(l)
doc.core_properties.title='Letter of the Day: routine script (English review copy)';doc.core_properties.language='en-US'
doc.add_heading('Letter of the Day: routine script (English review copy)',0)
doc.add_paragraph('This is an English translation of the script in the Spanish "Letra del día" app, written for review by a non-Spanish-speaking specialist. The Spanish script is the version teachers use. Text in quotation marks is what the teacher says aloud. Replace [letter name], [sound] and [word] with the values in the table at the end.')
doc.add_paragraph('Routine: Explicit Alphabet Knowledge Instructional Routine (Jones et al., 2013), adapted to Spanish. Each step has three rounds. In each round, change what you point to: the capital letter, the lowercase letter, or both.')
def sec(h,items):
    doc.add_heading(h,1)
    for kind,t in items:
        p=doc.add_paragraph(t,style='List Bullet' if kind=='b' else None)
        if kind=='s':
            p.paragraph_format.left_indent=Pt(18)
            for r in p.runs:r.italic=True
sec('1. Name it (default setting: letter names off, see Settings below)',[('s','"This is the letter [letter name]. This is the capital [letter name] and this is the lowercase [letter name]."'),('p','Point to each letter. Ask: "Which letter is this?" If the student does not respond, say the name and ask the student to repeat it or indicate it with their response method.'),('p','When letter names are turned off, the teacher says: "This is today\'s letter. This letter makes the sound [sound]." Bluebonnet\'s Kindergarten guide refers to letters by their sound, so names are off by default for Kindergarten.')])
sec('2. Its sound',[('s','"This letter stands for this sound: /[sound]/. Listen: /[sound]/ as in [word]."'),('p','Show the picture of the anchor word. Stretch the first sound of the word. Ask: "Which sound does this letter stand for?" The letter h has no sound. Say: "The h is written, but it is silent."'),('p','Letter x: "This letter can have two sounds, /s/ and /ks/. At the start of a word it sounds /s/, as in xilófono (xylophone). In the middle of a word it sounds /ks/, as in taxi." The app says /ks/ and then /s/. Source: Bluebonnet K Unit 5, Lesson 15, pp. 236 and 241 to 242.')])
sec('3. Find it',[('s','"Now let\'s find the letter [letter name]. When you find it, say the name and the sound."'),('p','The app reads one sentence aloud and highlights each word as it is read. A picture of the scene is shown with the sentence. The student then looks for the target letter in the sentence and taps it. When the student finds it, say the name and the sound. The large check mark appears on its own when every target letter has been found. Students are not asked to read the sentence.')])
sec('Syllables (consonants)',[('p','Point to a syllable in the row at the bottom of the board. Ask students to think of words that contain it. Write or point to their ideas in the box under the syllable. Tap the box to flip it and show two pictures of words that contain the syllable. Say the word and point to the highlighted syllable.'),('p','Tapping the syllable itself opens a larger window with the same two words and pictures, with Previous and Next buttons. Students listen and look at the syllable. They do not read the words.'),('p','For x, the syllable xi appears as /ks/ in taxi and as /s/ at the start of xilófono. The window shows both sounds.')])
sec('4. Write it',[('s','"I am going to show you how to write the letter [letter name]. First the capital, then the lowercase. As I write, I say each step."'),('p','Read the steps from the stroke guide while tracing. The student traces in the air, on paper, on a surface, or with an alternative pencil (keyboard, letter board, AAC device, eye gaze, adapted pencil). Show how to use that pencil. There are three checks for the capital letter and three for the lowercase letter.'),('p','Letter rr: rr appears only between vowels (perro, carro). It never starts a word, and the trilled sound at the start of a word is spelled with a single r (rana). For this reason the app shows rr as a lowercase card only, and the writing practice uses R and r, with the stroke steps from K Unit 3, Lesson 14.'),('p','Letter ch: the guide shows two letters. The first is the capital C and the second is the lowercase h (K Unit 5, Lesson 4, pp. 71 and 72).')])
sec('How the student can respond without speaking',[('b','Points to the letter, picture or choice with a finger, hand, pointer or head.'),('b','Looks at the choice when the cards are presented separately and held still.'),('b','Uses a communication device (AAC), a keyboard or a letter board.'),('b','If the student cannot point, use partner-assisted scanning: point to or read each choice aloud and the student answers yes or no with their own method.')])
doc.add_heading('Before use: specialist review',1)
for t in ['Speech-language pathologist (SLP) or AAC specialist: review before asking for isolated speech sounds from students with apraxia of speech, dysarthria or no functional speech.','Teacher of students with visual impairments (TVI): review size, contrast and tactile format of the cards and stroke guides.','Teacher of the deaf or audiologist: review before using the recorded sounds with students who have hearing loss. The recorded sounds are from Sounds of Speech (University of Iowa).','Anchor words marked with * were taken from the Bluebonnet lesson that teaches the letter. The others were chosen by the adapter and need to be verified. The sentences were written by the adapter, except three from K Unit 2, p. 311.']:
    doc.add_paragraph(t,style='List Bullet')
doc.add_heading('Settings',1)
for t in ['Use letter names: off by default. Bluebonnet K avoids letter names and refers to letters by their sound.','Theme: same as device, light, dark or high contrast.']:
    doc.add_paragraph(t,style='List Bullet')
doc.add_heading('Table by letter',1)
tbl=doc.add_table(rows=1,cols=6);tbl.style='Table Grid'
for c,h in zip(tbl.rows[0].cells,['Letter','Spanish letter name','Sound','Anchor word (English)','Sentence (Spanish)','Sentence (English)']):
    c.text=h
    for r in c.paragraphs[0].runs:r.bold=True
trPr=tbl.rows[0]._tr.get_or_add_trPr();th=OxmlElement('w:tblHeader');th.set(qn('w:val'),'true');trPr.append(th)
VERIF={'m','j','h','k','x','ch'}
for t in d['LET']:
    row=tbl.add_row().cells
    letter=t if t=='rr' else t[0].upper()+t[1:]+' '+t
    snd=SND.get(t,'/'+t+'/')
    vals=[letter,d['NAMES'][t],snd,'%s (%s)%s'%(d['ANCH'][t],AEN[t],' *' if t in VERIF else ''),L[t],EN[t]]
    for c,v in zip(row,vals):c.text=v
doc.add_paragraph()
doc.add_paragraph('Notes for the reviewer. Vowels have no syllable row. The sound column describes the target Spanish sound as used in Mexican and Texas Spanish. The anchor pictures come from the Bluebonnet picture cards where a card exists. Seven anchors have no Bluebonnet card (bata, chocolate, tiburón, wifi, xilófono, yoyo, zorro) and use Canva pictures instead.')
doc.add_paragraph('Sources: Explicit Alphabet Knowledge Instructional Routine (Jones et al., 2013), adapted to Spanish. Anchor words, stroke guides and text: Bluebonnet Learning, Texas Education Agency, CC BY-NC-SA 4.0. Isolated sounds: Sounds of Speech, University of Iowa. Sentence and word pictures: Canva.')
doc.save('../docs/Letter_of_the_Day_script_EN.docx')
